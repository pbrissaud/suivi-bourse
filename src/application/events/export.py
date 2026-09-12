"""The export: the ledger back out, in the format it comes in (issue #710)."""
import csv
import io
import re
import unicodedata
from datetime import date
from typing import (
    Any, Dict, Iterable, List, Mapping, NamedTuple, Optional, Sequence, Tuple,
)

from application import instants

from .loader import BASE_CURRENCY_COLUMN
from .schemas import DEFAULT_ACCOUNT, Event, unit_cost

EVENT_COLUMNS = (
    'date', 'event_type', 'account', 'symbol', 'name',
    'quantity', 'unit_price', 'fee', 'amount', 'notes',
    BASE_CURRENCY_COLUMN,
)

_UNWRITABLE = re.compile(r'[\x00-\x08\x0b\x0c\x0e-\x1f]')

UNDATED_SHEET = 'events'


def render_events(events: Iterable[Event],
                  base_currency: Optional[str] = None) -> str:
    """The ledger as an importable ``.csv``."""
    return _render(EVENT_COLUMNS, [
        _event_row(event, base_currency) for event in events
    ])


def _event_row(event: Event, base_currency: Optional[str]) -> dict:
    """One event as the cells a file holds."""
    return {
        'date': event.date,
        'event_type': event.event_type.value,
        'account': event.account,
        'symbol': event.symbol,
        'name': event.name,
        'quantity': event.quantity,
        'unit_price': event.unit_price,
        'fee': event.fee,
        'amount': event.amount,
        'notes': event.notes,
        BASE_CURRENCY_COLUMN: base_currency,
    }


def _render(columns: Sequence[str], rows: Sequence[dict]) -> str:
    """Header plus rows, ``\n``-terminated, UTF-8 without a byte-order mark."""
    buffer = io.StringIO(newline='')
    writer = csv.DictWriter(buffer, fieldnames=list(columns), restval='',
                            extrasaction='ignore', lineterminator='\n')
    writer.writeheader()
    for row in rows:
        writer.writerow({column: _cell(row.get(column)) for column in columns})
    return buffer.getvalue()


def _cell(value: Any) -> str:
    """One cell as text."""
    if value is None:
        return ''
    if isinstance(value, date):
        return instants.iso(value)
    return str(value)


def render_events_workbook(events: Iterable[Event],
                           base_currency: Optional[str] = None) -> bytes:
    """The ledger as an importable ``.xlsx``, **one sheet per year**."""
    import openpyxl

    workbook = openpyxl.Workbook()
    workbook.remove(workbook.active)
    for year, rows in _by_year(events).items():
        sheet = workbook.create_sheet(title=year)
        _write_row(sheet, list(EVENT_COLUMNS))
        for event in rows:
            cells = _event_row(event, base_currency)
            _write_row(sheet, [cells[column] for column in EVENT_COLUMNS])
    if not workbook.sheetnames:
        _write_row(workbook.create_sheet(title=UNDATED_SHEET),
                   list(EVENT_COLUMNS))

    buffer = io.BytesIO()
    workbook.save(buffer)
    workbook.close()
    return buffer.getvalue()


def _by_year(events: Iterable[Event]) -> Dict[str, List[Event]]:
    """The events grouped by the year they are dated, oldest tab first."""
    years: Dict[str, List[Event]] = {}
    for event in sorted(events, key=lambda event: event.date):
        years.setdefault(f'{event.date.year:04d}', []).append(event)
    return years


def _write_row(sheet, values: Sequence[Any]) -> None:
    """One row of cells, with text kept text and kept **writable**."""
    sheet.append([_writable(value) for value in values])
    for cell in sheet[sheet.max_row]:
        if isinstance(cell.value, str):
            cell.data_type = 's'


def _writable(value: Any) -> Any:
    """One value a worksheet will accept — text minus what OOXML cannot hold."""
    if not isinstance(value, str):
        return value
    return _UNWRITABLE.sub('', value)


class Selection(NamedTuple):
    """What the ledger's chips retain, as the export resource takes it."""

    query: str = ''
    event_type: Optional[str] = None
    account: Optional[str] = None
    symbols: Optional[Tuple[str, ...]] = None
    since: Optional[date] = None
    until: Optional[date] = None

    @property
    def reduces(self) -> bool:
        """Whether anything at all is being held back."""
        return bool(self.query.strip()) or self.event_type is not None \
            or self.account is not None or self.symbols is not None \
            or self.since is not None or self.until is not None


NO_SELECTION = Selection()


def select(events: Iterable[Event], selection: Selection) -> List[Event]:
    """The events a selection retains, in the order they arrived."""
    needle = fold(selection.query.strip())
    symbols = None if selection.symbols is None else set(selection.symbols)
    kept = []
    for event in events:
        if selection.since is not None and event.date < selection.since:
            continue
        if selection.until is not None and event.date > selection.until:
            continue
        if selection.event_type is not None \
                and event.event_type.value != selection.event_type:
            continue
        if selection.account is not None \
                and account_of(event) != selection.account:
            continue
        if symbols is not None \
                and (not event.symbol or event.symbol not in symbols):
            continue
        if needle and needle not in _haystack(event):
            continue
        kept.append(event)
    return kept


def account_of(event: Event) -> str:
    """The account a row names — **a blank one names** ``default``."""
    return (event.account or '').strip() or DEFAULT_ACCOUNT


def fold(value: str) -> str:
    """Accents dropped and case folded: a French label is searched as it is heard."""
    decomposed = unicodedata.normalize('NFD', value)
    return ''.join(character for character in decomposed
                   if not unicodedata.combining(character)).lower()


def _haystack(event: Event) -> str:
    """What the search reads: everything the identity and account columns show."""
    return fold(' '.join(part for part in
                         (event.symbol, event.notes, account_of(event)) if part))


# **No `account_type`** since #916 (ADR-0043): the column it rendered is written
# by the app and read by nobody, so exporting it would ship a field whose every
# row carries the same seeded word.
PORTFOLIO_COLUMNS = (
    'account', 'account_label',
    'cash_balance', 'net_contributed',
    'symbol', 'name', 'quantity', 'unit_cost', 'cost_basis',
    'price', 'market_value', 'realized_gain', 'received_dividend',
    BASE_CURRENCY_COLUMN,
)


def render_portfolio(accounts: Iterable[Any],
                     states: Mapping[str, Any],
                     positions: Iterable[Mapping[str, Any]],
                     base_currency: Optional[str] = None) -> str:
    """The accounts and their positions — balances, PMP and valuations (#836)."""
    declared = {account.id: account for account in accounts}
    held: Dict[str, List[Mapping[str, Any]]] = {}
    for position in positions:
        account = (position.get('account') or '').strip() or DEFAULT_ACCOUNT
        held.setdefault(account, []).append(position)

    rows: List[dict] = []
    for account in sorted(set(declared) | set(held)):
        rows.append(_account_row(account, declared.get(account),
                                 states.get(account), base_currency))
        for position in sorted(held.get(account, []),
                               key=lambda row: str(row.get('symbol') or '')):
            rows.append(_position_row(account, declared.get(account),
                                      position, base_currency))
    return _render(PORTFOLIO_COLUMNS, rows)


def _account_row(account: str, declaration: Optional[Any],
                 state: Optional[Any], base_currency: Optional[str]) -> dict:
    """The account itself: what it is called, and what cash stands in it."""
    return {
        'account': account,
        'account_label': None if declaration is None else declaration.label,
        'cash_balance': None if state is None else state.cash_balance,
        'net_contributed': None if state is None else state.net_contributed,
        BASE_CURRENCY_COLUMN: base_currency,
    }


def _position_row(account: str, declaration: Optional[Any],
                  position: Mapping[str, Any],
                  base_currency: Optional[str]) -> dict:
    """One holding: what is held, what it cost, and what it is worth."""
    quantity = position.get('quantity')
    price = position.get('price')
    return {
        'account': account,
        'account_label': None if declaration is None else declaration.label,
        'symbol': position.get('symbol'),
        'name': position.get('name'),
        'quantity': quantity,
        'unit_cost': unit_cost(quantity or 0.0,
                               position.get('cost_basis') or 0.0),
        'cost_basis': position.get('cost_basis'),
        'price': price,
        'market_value': (None if quantity is None or price is None
                         else quantity * price),
        'realized_gain': position.get('realized_gain'),
        'received_dividend': position.get('received_dividend'),
        BASE_CURRENCY_COLUMN: base_currency,
    }


__all__ = [
    'EVENT_COLUMNS', 'BASE_CURRENCY_COLUMN', 'PORTFOLIO_COLUMNS',
    'UNDATED_SHEET', 'NO_SELECTION', 'Selection',
    'render_events', 'render_events_workbook', 'render_portfolio',
    'select', 'account_of', 'fold',
]
